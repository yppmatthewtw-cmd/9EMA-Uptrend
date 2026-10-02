#!/usr/bin/env python3
"""Recalculate every formula in the workbook with pycel and compare the
總表 scores with the scanner's own values (the user's standard from earlier
rounds: the sheet must reproduce the screen, 0 差異 at the shown precision).

Usage: python3 scripts/verify_xlsx.py <xlsx> <screen json>
"""
import json, sys
from openpyxl import load_workbook
from pycel import ExcelCompiler

xlsx, js = sys.argv[1], sys.argv[2]
scr = json.load(open(js))
wb = load_workbook(xlsx)
xc = ExcelCompiler(filename=xlsx)

ws = wb["總表"]
hdr = {c.value: c.column_letter for c in ws[1]}
sym_col = hdr["代號"]
checks = {"趨勢分數": "s_trend", "回調質素": "s_pb", "底部結構": "s_hl", "三角形分數": "s_tri", "綜合分數": "score"}
by_sym = {r["sym"]: r for r in scr["rows"] + scr["tier2"]}
n_formula = 0; errors = []; skipped = []; worst = 0.0; worst_cell = ""
for sheet in wb.worksheets:
    for row in sheet.iter_rows():
        for c in row:
            if isinstance(c.value, str) and c.value.startswith("="):
                n_formula += 1
                try:
                    v = xc.evaluate(f"'{sheet.title}'!{c.coordinate}")
                except Exception as e:                      # noqa
                    if "is not implemented" in repr(e):       # pycel lacks the function (MEDIAN); Excel has it
                        skipped.append((sheet.title, c.coordinate)); continue
                    errors.append((sheet.title, c.coordinate, repr(e)[:120])); continue
                if isinstance(v, str) and v.startswith("#"):
                    errors.append((sheet.title, c.coordinate, v))
for r in range(2, ws.max_row + 1):
    sym = ws[f"{sym_col}{r}"].value
    if sym not in by_sym:
        continue
    for h, k in checks.items():
        v = xc.evaluate(f"'總表'!{hdr[h]}{r}")
        d = abs(float(v) - by_sym[sym][k])
        if d > worst:
            worst, worst_cell = d, f"{sym} {h}: sheet {v:.4f} vs scan {by_sym[sym][k]:.4f}"
print(f"formulas evaluated: {n_formula}; errors: {len(errors)}; skipped (function pycel lacks): {len(skipped)} {skipped[:6]}")
for e in errors[:20]:
    print("  ", e)
print(f"score check over {sum(1 for r in range(2, ws.max_row + 1) if ws[f'{sym_col}{r}'].value in by_sym)} rows x {len(checks)} scores: "
      f"max |sheet − scan| = {worst:.5f} ({worst_cell})")
sys.exit(1 if errors or worst > 0.05 else 0)
